
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

# =========================
# 1. 路径设置
# =========================
project_dir = Path(__file__).resolve().parent
output_dir = project_dir / "output"
chart_dir = output_dir / "charts"

rfm_path = output_dir / "customer_rfm_segmented.csv"
segment_path = output_dir / "customer_segment_summary.csv"
clean_path = output_dir / "online_retail_cleaned.csv"

report_path = output_dir / "Online_Retail_RFM_Analysis_Report.xlsx"

# 检查必需文件
for path in [rfm_path, segment_path, clean_path]:
    if not path.exists():
        raise FileNotFoundError(f"找不到文件：{path}")

rfm = pd.read_csv(rfm_path)
segments = pd.read_csv(segment_path)
clean = pd.read_csv(clean_path)

# =========================
# 2. 创建工作簿
# =========================
wb = Workbook()
ws_summary = wb.active
ws_summary.title = "Executive Summary"

ws_segments = wb.create_sheet("Segment Summary")
ws_repeat = wb.create_sheet("Repeat Purchase")
ws_rfm = wb.create_sheet("Customer RFM")
ws_quality = wb.create_sheet("Data Quality")
ws_charts = wb.create_sheet("Charts")

# 统一样式
header_fill = PatternFill(
    fill_type="solid", fgColor="1F4E78"
)
header_font = Font(
    color="FFFFFF", bold=True
)
title_font = Font(
    size=16, bold=True, color="1F1F1F"
)

def write_dataframe(ws, df):
    """将 DataFrame 写入工作表并设置表头。"""
    ws.append(list(df.columns))

    for row in df.itertuples(index=False, name=None):
        ws.append([
            None if pd.isna(value) else value
            for value in row
        ])

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )
        max_length = max(
            len(str(cell.value))
            if cell.value is not None else 0
            for cell in column_cells
        )
        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 30
        )

# =========================
# 3. Executive Summary
# =========================
total_customers = rfm["CustomerID"].nunique()
total_orders = clean["InvoiceNo"].nunique()
total_sales = clean["TotalAmount"].sum()

repeat_orders = (
    clean.groupby("CustomerID")["InvoiceNo"]
    .nunique()
)

repeat_customers = int((repeat_orders >= 2).sum())
one_time_customers = int((repeat_orders == 1).sum())

repeat_rate = repeat_customers / total_customers
one_time_rate = one_time_customers / total_customers

champions = segments.loc[
    segments["CustomerSegment"] == "Champions"
]

champion_sales = (
    float(champions["TotalMonetary"].iloc[0])
    if not champions.empty else 0
)
champion_share = (
    champion_sales / total_sales
    if total_sales else 0
)

ws_summary["A1"] = "Online Retail 电商客户复购与 RFM 分析报告"
ws_summary["A1"].font = title_font
ws_summary.merge_cells("A1:C1")

ws_summary.append([""])
ws_summary.append(["核心指标", "分析结果", "说明"])

metrics = [
    ["清洗后交易记录数", len(clean), "剔除缺失客户、取消订单及无效数量/价格记录"],
    ["客户总数", total_customers, "按 CustomerID 去重"],
    ["有效订单数", total_orders, "按 InvoiceNo 去重"],
    ["累计销售额", round(total_sales, 2), "Quantity × UnitPrice 汇总"],
    ["复购客户数", repeat_customers, "至少有 2 个不同订单"],
    ["复购客户占比", repeat_rate, "复购客户数 / 客户总数"],
    ["一次购买客户数", one_time_customers, "只有 1 个不同订单"],
    ["一次购买客户占比", one_time_rate, "一次购买客户数 / 客户总数"],
    ["Champions 客户数", int(champions["CustomerCount"].sum()), "根据当前 RFM 业务规则分类"],
    ["Champions 销售额", round(champion_sales, 2), "该分层客户的累计消费额"],
    ["Champions 销售额占比", champion_share, "Champions 销售额 / 全部销售额"],
]

for row in metrics:
    ws_summary.append(row)

for cell in ws_summary[3]:
    cell.fill = header_fill
    cell.font = header_font

for row in range(4, ws_summary.max_row + 1):
    label = ws_summary.cell(row, 1).value
    value_cell = ws_summary.cell(row, 2)

    if "占比" in str(label):
        value_cell.number_format = "0.00%"
    elif "销售额" in str(label):
        value_cell.number_format = '#,##0.00'

ws_summary.column_dimensions["A"].width = 28
ws_summary.column_dimensions["B"].width = 20
ws_summary.column_dimensions["C"].width = 62
ws_summary.freeze_panes = "A4"

ws_summary.append([""])
ws_summary.append(["主要发现"])
ws_summary.cell(ws_summary.max_row, 1).font = Font(bold=True, size=12)

findings = [
    f"清洗后共有 {total_customers:,} 名客户，产生 {total_orders:,} 个有效订单。",
    f"复购客户占比为 {repeat_rate:.2%}，可用于描述该数据集中的客户复购情况。",
    f"Champions 客户贡献约 {champion_share:.2%} 的累计销售额。",
    "沉睡客户人数较多，可将客户召回作为后续运营策略的测试方向。",
    "高价值流失风险客户可作为定向挽回的候选群体；实际效果仍需实验验证。",
]

for finding in findings:
    ws_summary.append([finding])
    ws_summary.cell(ws_summary.max_row, 1).alignment = Alignment(
        wrap_text=True, vertical="top"
    )
    ws_summary.merge_cells(
        start_row=ws_summary.max_row,
        start_column=1,
        end_row=ws_summary.max_row,
        end_column=3
    )

# =========================
# 4. 客户分层汇总
# =========================
write_dataframe(ws_segments, segments)

# 占比和金额格式
for row in range(2, ws_segments.max_row + 1):
    for col in range(1, ws_segments.max_column + 1):
        header = ws_segments.cell(1, col).value
        cell = ws_segments.cell(row, col)

        if "SharePct" in str(header):
            cell.number_format = "0.00"
        elif any(
            key in str(header)
            for key in ["Monetary", "TotalMonetary"]
        ):
            cell.number_format = '#,##0.00'

# =========================
# 5. 复购分析
# =========================
repeat_summary = pd.DataFrame([
    {
        "CustomerType": "Repeat Customers",
        "CustomerCount": repeat_customers,
        "CustomerSharePct": round(repeat_rate * 100, 2)
    },
    {
        "CustomerType": "One-time Buyers",
        "CustomerCount": one_time_customers,
        "CustomerSharePct": round(one_time_rate * 100, 2)
    }
])

write_dataframe(ws_repeat, repeat_summary)

ws_repeat["E1"] = "分析口径"
ws_repeat["E1"].font = header_font
ws_repeat["E1"].fill = header_fill
ws_repeat["E2"] = "复购客户：至少有 2 个不同订单的客户。"
ws_repeat["E3"] = "一次购买客户：只有 1 个不同订单的客户。"
ws_repeat["E4"] = "订单数按不同 InvoiceNo 计算。"
ws_repeat.column_dimensions["E"].width = 48

# =========================
# 6. 客户 RFM 明细
# =========================
write_dataframe(ws_rfm, rfm)

# =========================
# 7. 数据质量检查
# =========================
quality_rows = [
    ["检查项目", "清洗前/原始数据", "清洗后数据", "检查说明"],
    ["记录数", 541909, len(clean), "保留有效购买记录"],
    ["CustomerID 缺失数", 135080, int(clean["CustomerID"].isna().sum()), "清洗后应为 0"],
    ["取消订单数量（发票数以 C 开头的记录）", 9288, int(clean["InvoiceNo"].astype(str).str.startswith("C").sum()), "清洗后应为 0"],
    ["Quantity <= 0 记录数", 10624, int((clean["Quantity"] <= 0).sum()), "清洗后应为 0"],
    ["UnitPrice <= 0 记录数", 2517, int((clean["UnitPrice"] <= 0).sum()), "清洗后应为 0"],
    ["清洗后客户数", "-", clean["CustomerID"].nunique(), "与 RFM 客户数核对"],
    ["清洗后订单数", "-", clean["InvoiceNo"].nunique(), "按 InvoiceNo 去重"],
    ["清洗后销售额", "-", round(clean["TotalAmount"].sum(), 2), "与 TotalAmount 求和核对"],
]

for row in quality_rows:
    ws_quality.append(row)

for cell in ws_quality[1]:
    cell.fill = header_fill
    cell.font = header_font

ws_quality.freeze_panes = "A2"
ws_quality.auto_filter.ref = ws_quality.dimensions

for col, width in {"A": 38, "B": 20, "C": 20, "D": 38}.items():
    ws_quality.column_dimensions[col].width = width

# =========================
# 8. 插入三张图表
# =========================
chart_files = [
    ("01_customer_segment_count.png", "客户分层人数"),
    ("02_customer_segment_sales.png", "客户分层销售额"),
    ("03_frequency_vs_monetary.png", "购买频率与消费金额"),
]

ws_charts["A1"] = "Customer Segmentation Visualizations"
ws_charts["A1"].font = title_font

row_position = 3

for filename, title in chart_files:
    image_path = chart_dir / filename

    if image_path.exists():
        ws_charts.cell(row_position, 1, title)
        ws_charts.cell(row_position, 1).font = Font(
            bold=True, size=12
        )

        img = XLImage(str(image_path))
        img.width = 720
        img.height = 432
        ws_charts.add_image(img, f"A{row_position + 1}")

        row_position += 25
    else:
        ws_charts.cell(
            row_position, 1, f"未找到图表文件：{filename}"
        )
        row_position += 2

# =========================
# 9. 保存报告
# =========================
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

wb.save(report_path)

print("Excel 分析报告生成成功！")
print("保存位置：", report_path)
print("客户数：", total_customers)
print("订单数：", total_orders)
print("复购客户占比：", f"{repeat_rate:.2%}")
print("Champions 销售额占比：", f"{champion_share:.2%}")
